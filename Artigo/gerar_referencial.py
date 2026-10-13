# -*- coding: utf-8 -*-
"""Gera 'Referencial Teorico.docx' seguindo a ABNT (NBR 6022:2018, 10520:2023, 6023:2018)."""
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"

doc = Document()

# ---------- Pagina e margens (ABNT: sup/esq 3 cm; inf/dir 2 cm) ----------
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.orientation = WD_ORIENT.PORTRAIT
sec.top_margin, sec.left_margin = Cm(3), Cm(3)
sec.bottom_margin, sec.right_margin = Cm(2), Cm(2)
sec.header_distance = Cm(1.25)

# ---------- Estilo base ----------
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.5, indent=True,
        before=0, after=0, left=None, keep_next=False):
    pf = p.paragraph_format
    p.alignment = align
    pf.line_spacing = spacing
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.first_line_indent = Cm(1.25) if indent else None
    if left is not None:
        pf.left_indent = Cm(left)
    pf.keep_with_next = keep_next


def runs(p, text, size=12, bold_all=False):
    """Suporta **negrito** e *italico* inline."""
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*"):
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            r = p.add_run(tok)
        r.font.name, r.font.size = FONT, Pt(size)
        if bold_all:
            r.bold = True


def para(text):
    p = doc.add_paragraph()
    fmt(p)
    runs(p, text)


def h1(text):  # secao primaria: MAIUSCULAS, negrito, sem recuo
    p = doc.add_paragraph()
    fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=12, after=12, keep_next=True)
    runs(p, text.upper(), bold_all=True)


def h2(text):  # secao secundaria: minuscula, negrito
    p = doc.add_paragraph()
    fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=12, after=12, keep_next=True)
    runs(p, text, bold_all=True)


# ---------- Numero de pagina (canto superior direito) ----------
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run()
r.font.name, r.font.size = FONT, Pt(10)
for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if t:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), t)
    else:
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
    r._r.append(el)

# ---------- Titulo provisorio ----------
p = doc.add_paragraph()
fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=18)
runs(p, "INTELIGÊNCIA ARTIFICIAL GENERATIVA NO ENSINO DE PROGRAMAÇÃO: PENSAMENTO "
        "COMPUTACIONAL, AUTOEFICÁCIA E LETRAMENTO EM IA NO ENSINO MÉDIO", bold_all=True)

# =====================================================================
h1("2 Referencial teórico")
para("Este capítulo apresenta os fundamentos que sustentam a investigação, organizados em "
     "seis eixos: (i) pensamento computacional e seu ensino na educação básica; (ii) instrumentos "
     "de avaliação do pensamento computacional; (iii) autoeficácia em programação; (iv) inteligência "
     "artificial (IA) generativa no ensino de programação; (v) letramento em IA e engenharia de "
     "*prompt*; e (vi) estratégias pedagógicas de ensino de programação para iniciantes.")

# ---------------------------------------------------------------------
h2("2.1 Pensamento computacional e seu ensino na educação básica")
para("O termo pensamento computacional ganhou projeção com Wing (2006), que o descreve como o "
     "conjunto de processos de pensamento envolvidos na resolução de problemas, na concepção de "
     "sistemas e na compreensão do comportamento humano com base em conceitos fundamentais da "
     "ciência da computação, defendendo-o como uma habilidade a ser desenvolvida por todos, e não "
     "apenas por profissionais da área. Em uma perspectiva mais próxima da prática escolar, "
     "Brennan e Resnick (2012) organizam o pensamento computacional em três dimensões: conceitos "
     "computacionais (como sequências, laços, condicionais e eventos), práticas computacionais "
     "(como experimentar, testar, depurar e reutilizar) e perspectivas computacionais (a forma "
     "como o indivíduo se relaciona com o mundo e consigo mesmo por meio da computação).")
para("A aprendizagem por meio da construção de artefatos computacionais remonta a Papert (1980), "
     "para quem programar permite ao estudante explorar ideias e corrigir seus próprios erros de "
     "modo ativo, o que fundamenta abordagens construcionistas. Grover e Pea (2013), ao revisarem "
     "o estado da área na educação básica (K-12), destacam a importância de definir de forma "
     "operacional o que se entende por pensamento computacional e de ampliar as pesquisas sobre "
     "como avaliá-lo.")
para("No Brasil, a Resolução CNE/CEB nº 1, de 4 de outubro de 2022, instituiu normas sobre "
     "computação na educação básica como complemento à Base Nacional Comum Curricular (BNCC), "
     "organizando o conteúdo em torno dos eixos pensamento computacional, mundo digital e cultura "
     "digital (BRASIL, 2022). Nesse contexto, Brackmann (2017) investigou o desenvolvimento do "
     "pensamento computacional na educação básica brasileira e contribuiu para a adaptação e o uso "
     "de instrumentos de avaliação no país, constituindo referência nacional para estudos da área.")

# ---------------------------------------------------------------------
h2("2.2 Avaliação do pensamento computacional: Bebras e Teste de Pensamento Computacional")
para("A mensuração do pensamento computacional é um desafio recorrente na literatura. Entre as "
     "iniciativas de ampla adoção está o Bebras, desafio internacional composto por tarefas "
     "curtas que avaliam raciocínio algorítmico e lógico sem exigir conhecimento prévio de "
     "programação (DAGIENĖ; STUPURIENĖ, 2016). Essa característica o torna adequado para "
     "públicos iniciantes e para medições realizadas antes do contato dos estudantes com a "
     "linguagem de programação.")
para("Outro instrumento consolidado é o *Computational Thinking Test* (CTt), proposto por "
     "Román-González, Pérez-González e Jiménez-Fernández (2017) com o objetivo de verificar quais "
     "habilidades cognitivas se relacionam ao pensamento computacional, e concebido para "
     "estudantes adolescentes. Delineamentos que aplicam instrumentos desse tipo em dois momentos "
     "(pré-teste e pós-teste) permitem estimar a variação do desempenho dos participantes após "
     "uma intervenção, o que é pertinente para investigar o efeito de uma trilha formativa que "
     "integra programação e IA.")

# ---------------------------------------------------------------------
h2("2.3 Autoeficácia em programação")
para("A autoeficácia, conceito central da teoria social cognitiva, refere-se à crença do "
     "indivíduo em sua capacidade de organizar e executar as ações necessárias para alcançar um "
     "resultado (BANDURA, 1977). Segundo o autor, essas crenças são construídas principalmente "
     "por quatro fontes: experiências de êxito, experiências vicárias (observação de outros), "
     "persuasão verbal e estados fisiológicos e emocionais. No ensino de programação, em que "
     "estudantes iniciantes frequentemente se deparam com mensagens de erro e frustração, "
     "estruturar atividades que proporcionem experiências graduais de êxito tende a ser "
     "particularmente relevante à luz dessa teoria.")
para("Para mensurar esse construto no contexto da programação, Tsai, Wang e Hsu (2019) "
     "desenvolveram a *Computer Programming Self-Efficacy Scale* (CPSES), escala multidimensional "
     "respondida em formato Likert, que vem sendo utilizada para acompanhar a confiança de "
     "estudantes em atividades de programação. Sua aplicação em momentos distintos possibilita "
     "analisar se uma intervenção pedagógica se associa a mudanças na percepção de competência "
     "dos participantes.")

# ---------------------------------------------------------------------
h2("2.4 Inteligência artificial generativa no ensino de programação")
para("A disseminação de modelos de linguagem de grande porte capazes de gerar, explicar e "
     "corrigir código alterou o cenário do ensino de computação. Becker *et al.* (2023) discutem "
     "que a geração automática de código apresenta, ao mesmo tempo, oportunidades (como "
     "exemplos personalizados, explicações e apoio à depuração) e desafios (como o uso indevido "
     "para obter respostas prontas e a necessidade de repensar avaliações). Denny *et al.* (2024) "
     "reforçam que tais modelos resolvem com facilidade problemas típicos de disciplinas "
     "introdutórias, o que exige a adaptação de objetivos de aprendizagem e de práticas "
     "pedagógicas, em vez da simples proibição da tecnologia.")
para("As evidências empíricas sobre o efeito dessas ferramentas em iniciantes são, contudo, "
     "matizadas. Kazemitabaar *et al.* (2023) observaram que aprendizes novatos que dispunham de "
     "um gerador de código obtiveram melhor desempenho em tarefas de autoria, sem prejuízo "
     "evidente em avaliações posteriores realizadas sem o apoio da ferramenta. Por outro lado, "
     "Prather *et al.* (2024) alertam para o risco de ampliação das desigualdades entre os "
     "estudantes: aqueles com menor capacidade metacognitiva podem ser induzidos a uma ilusão de "
     "competência, enquanto os mais autorregulados tendem a aproveitar melhor o recurso. Esses "
     "resultados indicam que o ganho de aprendizagem depende do modo como a IA é inserida no "
     "processo didático, o que justifica pesquisas que descrevam e avaliem propostas "
     "estruturadas de uso orientado.")

# ---------------------------------------------------------------------
h2("2.5 Letramento em IA e engenharia de *prompt*")
para("O letramento em IA é definido por Long e Magerko (2020) como o conjunto de competências "
     "que permite às pessoas avaliar criticamente tecnologias de IA, comunicar-se e colaborar "
     "com elas de forma eficaz e utilizá-las como ferramenta em diferentes contextos. Esse "
     "conceito amplia o foco do simples uso instrumental para a compreensão dos limites, dos "
     "riscos e das implicações éticas desses sistemas, dimensões importantes na formação de "
     "adolescentes.")
para("Nesse cenário, a engenharia de *prompt* emerge como uma habilidade de comunicação com os "
     "modelos de linguagem. White *et al.* (2023) propõem um catálogo de padrões de *prompt*, "
     "entendidos como soluções reutilizáveis para problemas recorrentes na interação com "
     "modelos como o ChatGPT, o que evidencia que a qualidade da instrução formulada influencia "
     "a qualidade da resposta obtida. No ensino de programação, formular um *prompt* claro "
     "exige que o estudante decomponha o problema, explicite requisitos e restrições e avalie "
     "criticamente o código recebido, habilidades associadas ao próprio pensamento computacional "
     "(BRENNAN; RESNICK, 2012).")

# ---------------------------------------------------------------------
h2("2.6 Estratégias pedagógicas para o ensino de programação a iniciantes")
para("Kirschner, Sweller e Clark (2006) argumentam que, para aprendizes iniciantes, "
     "abordagens com mínima orientação são menos eficazes do que aquelas com instrução "
     "guiada, o que respalda sequências didáticas estruturadas com demonstração, prática "
     "orientada e retirada gradual do apoio. Essa lógica dialoga com o conceito de zona de "
     "desenvolvimento proximal de Vygotsky (1991), segundo o qual o apoio de um mediador mais "
     "experiente permite ao estudante realizar tarefas que ainda não realizaria sozinho.")
para("Estudos multinacionais indicam que muitos programadores novatos apresentam dificuldades "
     "para ler e rastrear a execução de trechos de código (LISTER *et al.*, 2004). Em resposta, "
     "Sentance, Waite e Kallia (2019) propõem a abordagem PRIMM (*Predict, Run, Investigate, "
     "Modify, Make*), que antecipa a leitura, a previsão e a investigação de código existente "
     "antes da escrita autônoma. Estratégias de análise de código pronto e de depuração "
     "orientada, como a engenharia reversa guiada e os testes de estresse, são coerentes com "
     "esse princípio.")
para("Por fim, a aprendizagem baseada em projetos favorece o engajamento ao propor problemas "
     "autênticos e a produção de artefatos concretos (BLUMENFELD *et al.*, 1991). Quando "
     "vinculados aos Objetivos de Desenvolvimento Sustentável da Agenda 2030 (ORGANIZAÇÃO DAS "
     "NAÇÕES UNIDAS, 2015), os projetos agregam à dimensão técnica uma dimensão cidadã, "
     "estimulando a aplicação da programação a problemas socialmente relevantes.")

# ---------------------------------------------------------------------
h2("2.7 Síntese e lacuna de pesquisa")
para("O Quadro 1 sintetiza a relação entre os construtos investigados, os instrumentos de "
     "medida e os principais referenciais adotados.")

# Quadro 1 (titulo acima, fonte abaixo - NBR 14724)
p = doc.add_paragraph(); fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.0,
                              before=6, after=6, keep_next=True)
runs(p, "Quadro 1 – Construtos, instrumentos e referenciais teóricos", size=10)

rows = [
    ("Construto", "Instrumento / foco", "Referenciais"),
    ("Pensamento computacional", "Bebras e CTt (pré e pós-teste)",
     "Wing (2006); Brennan e Resnick (2012); Dagienė e Stupurienė (2016); "
     "Román-González, Pérez-González e Jiménez-Fernández (2017)"),
    ("Autoeficácia em programação", "CPSES (escala Likert)", "Bandura (1977); Tsai, Wang e Hsu (2019)"),
    ("Letramento em IA", "Questionário de letramento em IA; análise de *prompts*",
     "Long e Magerko (2020); White *et al.* (2023)"),
    ("Uso de IA generativa na aprendizagem", "Registros de interação com a IA e produção de código",
     "Becker *et al.* (2023); Kazemitabaar *et al.* (2023); Prather *et al.* (2024)"),
]
tbl = doc.add_table(rows=len(rows), cols=3)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, row in enumerate(rows):
    for j, txt in enumerate(row):
        c = tbl.cell(i, j)
        c.paragraphs[0].text = ""
        pp = c.paragraphs[0]
        fmt(pp, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, spacing=1.0)
        runs(pp, txt, size=10, bold_all=(i == 0))

p = doc.add_paragraph(); fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, spacing=1.0,
                              before=4, after=12)
runs(p, "Fonte: Elaborado pelos autores (2026).", size=10)

para("Observa-se, portanto, que a literatura já discute as oportunidades e os riscos da IA "
     "generativa para iniciantes e dispõe de instrumentos validados para mensurar pensamento "
     "computacional e autoeficácia. Entretanto, ainda são escassos os estudos que descrevam e "
     "avaliem, de forma integrada, uma trilha formativa para estudantes do ensino médio sem "
     "conhecimento prévio em programação, que combine instrução guiada, engenharia de *prompt* "
     "e letramento em IA, mensurando conjuntamente seus efeitos sobre o pensamento "
     "computacional e a autoeficácia. É nessa lacuna que se insere a presente investigação.")

# =====================================================================
# REFERENCIAS (NBR 6023:2018) - espacamento simples, alinhadas a esquerda
doc.add_page_break()
p = doc.add_paragraph()
fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=12)
runs(p, "REFERÊNCIAS", bold_all=True)

REFS = [
    "BANDURA, Albert. Self-efficacy: toward a unifying theory of behavioral change. "
    "**Psychological Review**, Washington, v. 84, n. 2, p. 191-215, 1977.",

    "BECKER, Brett A. *et al.* Programming is hard – or at least it used to be: educational "
    "opportunities and challenges of AI code generation. In: ACM TECHNICAL SYMPOSIUM ON COMPUTER "
    "SCIENCE EDUCATION, 54., 2023, Toronto. **Proceedings** [...]. New York: ACM, 2023. p. 500-506.",

    "BLUMENFELD, Phyllis C. *et al.* Motivating project-based learning: sustaining the doing, "
    "supporting the learning. **Educational Psychologist**, Abingdon, v. 26, n. 3-4, p. 369-398, 1991.",

    "BRACKMANN, Christian Puhlmann. **Desenvolvimento do pensamento computacional através de "
    "atividades desplugadas na educação básica**. 2017. Tese (Doutorado em Informática na "
    "Educação) – Faculdade de Educação, Universidade Federal do Rio Grande do Sul, Porto Alegre, 2017.",

    "BRASIL. Conselho Nacional de Educação. Câmara de Educação Básica. **Resolução CNE/CEB nº 1, "
    "de 4 de outubro de 2022**. Institui Normas sobre Computação na Educação Básica – Complemento "
    "à Base Nacional Comum Curricular (BNCC). Brasília, DF: CNE, 2022.",

    "BRENNAN, Karen; RESNICK, Mitchel. New frameworks for studying and assessing the development "
    "of computational thinking. In: ANNUAL MEETING OF THE AMERICAN EDUCATIONAL RESEARCH "
    "ASSOCIATION, 2012, Vancouver. **Proceedings** [...]. Vancouver: AERA, 2012. p. 1-25.",

    "DAGIENĖ, Valentina; STUPURIENĖ, Gabrielė. Bebras – a sustainable community building model "
    "for the concept based learning of informatics and computational thinking. **Informatics in "
    "Education**, Vilnius, v. 15, n. 1, p. 25-44, 2016.",

    "DENNY, Paul *et al.* Computing education in the era of generative AI. **Communications of "
    "the ACM**, New York, v. 67, n. 2, p. 56-67, 2024.",

    "GROVER, Shuchi; PEA, Roy. Computational thinking in K–12: a review of the state of the "
    "field. **Educational Researcher**, Thousand Oaks, v. 42, n. 1, p. 38-43, 2013.",

    "KAZEMITABAAR, Majeed *et al.* Studying the effect of AI code generators on supporting novice "
    "learners in introductory programming. In: CHI CONFERENCE ON HUMAN FACTORS IN COMPUTING "
    "SYSTEMS, 2023, Hamburg. **Proceedings** [...]. New York: ACM, 2023. p. 1-23.",

    "KIRSCHNER, Paul A.; SWELLER, John; CLARK, Richard E. Why minimal guidance during instruction "
    "does not work: an analysis of the failure of constructivist, discovery, problem-based, "
    "experiential, and inquiry-based teaching. **Educational Psychologist**, Abingdon, v. 41, "
    "n. 2, p. 75-86, 2006.",

    "LISTER, Raymond *et al.* A multi-national study of reading and tracing skills in novice "
    "programmers. **ACM SIGCSE Bulletin**, New York, v. 36, n. 4, p. 119-150, 2004.",

    "LONG, Duri; MAGERKO, Brian. What is AI literacy? Competencies and design considerations. In: "
    "CHI CONFERENCE ON HUMAN FACTORS IN COMPUTING SYSTEMS, 2020, Honolulu. **Proceedings** [...]. "
    "New York: ACM, 2020. p. 1-16.",

    "ORGANIZAÇÃO DAS NAÇÕES UNIDAS. **Transformando nosso mundo**: a Agenda 2030 para o "
    "Desenvolvimento Sustentável. Rio de Janeiro: ONU Brasil, 2015.",

    "PAPERT, Seymour. **Mindstorms**: children, computers, and powerful ideas. New York: Basic "
    "Books, 1980.",

    "PRATHER, James *et al.* The widening gap: the benefits and harms of generative AI for novice "
    "programmers. In: ACM CONFERENCE ON INTERNATIONAL COMPUTING EDUCATION RESEARCH, 2024, "
    "Melbourne. **Proceedings** [...]. New York: ACM, 2024.",

    "ROMÁN-GONZÁLEZ, Marcos; PÉREZ-GONZÁLEZ, Juan-Carlos; JIMÉNEZ-FERNÁNDEZ, Carmen. Which "
    "cognitive abilities underlie computational thinking? Criterion validity of the Computational "
    "Thinking Test. **Computers in Human Behavior**, Amsterdam, v. 72, p. 678-691, 2017.",

    "SENTANCE, Sue; WAITE, Jane; KALLIA, Maria. Teaching computer programming with PRIMM: a "
    "sociocultural perspective. **Computer Science Education**, Abingdon, v. 29, n. 2-3, "
    "p. 136-176, 2019.",

    "TSAI, Meng-Jung; WANG, Chin-Yeh; HSU, Po-Fen. Developing the computer programming "
    "self-efficacy scale for computer literacy education. **Journal of Educational Computing "
    "Research**, Thousand Oaks, v. 56, n. 8, p. 1345-1360, 2019.",

    "VYGOTSKY, Lev Semenovich. **A formação social da mente**: o desenvolvimento dos processos "
    "psicológicos superiores. 4. ed. São Paulo: Martins Fontes, 1991.",

    "WHITE, Jules *et al.* **A prompt pattern catalog to enhance prompt engineering with "
    "ChatGPT**. [S. l.]: arXiv, 2023. Disponível em: https://arxiv.org/abs/2302.11382. Acesso em: "
    "13 out. 2026.",

    "WING, Jeannette M. Computational thinking. **Communications of the ACM**, New York, v. 49, "
    "n. 3, p. 33-35, 2006.",
]
for ref in REFS:
    p = doc.add_paragraph()
    fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, spacing=1.0, after=12)
    runs(p, ref)

out = r"c:\Users\Victor\OneDrive\PESSOAL\PROJETO FAPEMING\Artigo\Referencial Teorico.docx"
doc.save(out)
print("OK:", out)
